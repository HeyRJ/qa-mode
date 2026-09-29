#!/usr/bin/env python3
"""Trust, then check: read Claude's test case file and check it against the requirements.

usage: python3 check_testcases.py "<test cases>.xlsx" ["<requirements>.xlsx"] ["<open questions>.xlsx"]
       The requirements default to the Reach example's. The open questions file defaults to any
       "Open questions*.xlsx" next to the test cases. Needs openpyxl (pip install openpyxl).

The requirements file uses BA mode's layout: a "Stories" sheet with one row per acceptance criterion
(Story ID, Epic, Story, AC ID, Given, When, Then, Priority, ...).

Checks, each printed as ok / !! with details:
- the sheets are there (Test cases, Coverage, Test data; the open questions in their own file or sheet)
  and their header rows are frozen
- every AC ID in the requirements is covered by at least one test case (read from the test cases' AC column,
  not just from Coverage), and Coverage agrees with the test cases
- test case IDs are unique; every Status is "Not run"; every test case has steps, test data and an expected result
- Won't stories only check that the feature isn't there
- no vague expected results ("correct", "properly", "an error is shown", ...)
- the open questions file: its columns, unique IDs, every Status Open and every Answer empty, and it names only
  test cases and ACs that exist
Reach example only (KEYS and MSGS below: change them for your project, or delete that block):
- the sample file's key numbers appear in the expected results (10.5% A+, 11.1% A+, 2.6% C, row 7 at 3.0% B, ...)
- the seven "Rows to fix" messages of RC-04.2 appear word for word
- a case sits exactly on a rounding half (x.x5%)
Then counts by type and priority."""
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
tc_file = Path(sys.argv[1])
req_file = Path(sys.argv[2]) if len(sys.argv) > 2 else HERE / "example-reach" / "Reach report card v1 requirements.xlsx"

OK, BAD = "ok", "!!"
problems = 0


def say(flag, msg):
    global problems
    problems += flag == BAD
    print(f"{flag} {msg}")


def norm(s):
    return re.sub(r"[^a-z0-9]+", " ", str(s or "").lower()).strip()


def find_sheet(wb, *names):
    want = [norm(n) for n in names]
    for ws in wb.worksheets:
        n = norm(ws.title)
        if n in want or any(w in n for w in want):
            return ws
    return None


def header_map(ws):
    hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    return {norm(h): i for i, h in enumerate(hdr) if h is not None}, hdr


def col(hmap, *cands):
    for c in cands:
        c = norm(c)
        for k, i in hmap.items():
            if k == c:
                return i
    for c in cands:
        c = norm(c)
        for k, i in hmap.items():
            if c in k:
                return i
    return None


# ---------- the requirements ----------
req = openpyxl.load_workbook(req_file, read_only=True)
st = req["Stories"]
acs, prio_of, story_of = [], {}, {}
for r in st.iter_rows(min_row=2, values_only=True):
    if not r or not r[3]:
        continue
    acs.append(r[3]); prio_of[r[3]] = r[7]; story_of[r[3]] = r[0]
wont_stories = sorted({story_of[a] for a in acs if prio_of[a] == "Won't"})
print(f"requirements: {len(acs)} ACs in {len(set(story_of.values()))} stories "
      f"({dict(Counter(prio_of.values()))}); Won't stories: {', '.join(wont_stories)}")

# ---------- Claude's file ----------
wb = openpyxl.load_workbook(tc_file)
sheets = {k: find_sheet(wb, *v) for k, v in {
    "tc": ("Test cases", "Testcases", "Test case"), "cov": ("Coverage", "Traceability"),
    "data": ("Test data",), "open": ("Open questions", "Questions")}.items()}
oq_file = Path(sys.argv[3]) if len(sys.argv) > 3 else next(iter(sorted(tc_file.parent.glob("Open questions*.xlsx"))), None)
for k, name in [("tc", "Test cases"), ("cov", "Coverage"), ("data", "Test data"), ("open", "Open questions")]:
    ws = sheets[k]
    if ws is None and k == "open" and oq_file is not None:
        say(OK, f"no Open questions sheet; the open questions are in '{oq_file.name}' (checked below)")
    elif ws is None:
        say(BAD, f"sheet '{name}' is missing (sheets: {', '.join(wb.sheetnames)})")
    else:
        fz = ws.freeze_panes
        say(OK if fz and fz != "A1" else BAD, f"sheet '{ws.title}': {ws.max_row - 1} rows, header frozen: {fz or 'no'}")
ws = sheets["tc"]
if ws is None:
    sys.exit(1)
hm, hdr = header_map(ws)
c_id, c_story, c_ac = col(hm, "TC ID", "ID"), col(hm, "Story", "Story ID"), col(hm, "AC", "AC ID", "AC IDs")
c_type, c_title = col(hm, "Type"), col(hm, "Title")
c_pre, c_data, c_steps = col(hm, "Preconditions"), col(hm, "Test data"), col(hm, "Steps")
c_exp, c_prio, c_rules, c_status = col(hm, "Expected result", "Expected"), col(hm, "Priority"), col(hm, "Rules", "Rule IDs"), col(hm, "Status")
missing_cols = [n for n, c in [("TC ID", c_id), ("AC", c_ac), ("Type", c_type), ("Steps", c_steps), ("Expected result", c_exp),
                               ("Test data", c_data), ("Priority", c_prio), ("Status", c_status)] if c is None]
say(BAD if missing_cols else OK, f"test case columns: {', '.join(str(h) for h in hdr if h)}" +
    (f"; missing: {', '.join(missing_cols)}" if missing_cols else ""))

rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r and any(v not in (None, "") for v in r)]
get = lambda r, c: "" if c is None or c >= len(r) or r[c] is None else str(r[c]).strip()
ids = [get(r, c_id) for r in rows]
dup = [i for i, n in Counter(ids).items() if n > 1]
say(BAD if dup else OK, f"{len(rows)} test cases, IDs unique" + (f"; duplicates: {dup}" if dup else ""))

ac_re = re.compile(r"[A-Z]{2,}-\d{2}\.\d+")  # AC IDs like RC-01.1
covered = defaultdict(list)
for r in rows:
    for a in ac_re.findall(get(r, c_ac)):
        covered[a].append(get(r, c_id))
uncovered = [a for a in acs if a not in covered]
say(BAD if uncovered else OK, f"ACs covered by test cases: {len(acs) - len(uncovered)} of {len(acs)}" +
    (f"; not covered: {', '.join(uncovered)}" if uncovered else ""))
unknown = sorted(set(covered) - set(acs))
if unknown:
    say(BAD, f"test cases name AC IDs that aren't in the requirements: {', '.join(unknown)}")

# Coverage sheet agrees?
cws = sheets["cov"]
if cws is not None:
    chm, _ = header_map(cws)
    cc_ac, cc_tc = col(chm, "AC ID", "AC"), col(chm, "TC IDs", "TC ID", "Test cases", "TCs")
    cov = {}
    for r in cws.iter_rows(min_row=2, values_only=True):
        if not r or cc_ac is None:
            continue
        a = get(r, cc_ac)
        if ac_re.fullmatch(a):
            cov[a] = set(re.findall(r"TC-\d+", get(r, cc_tc))) if cc_tc is not None else set()
    miss = [a for a in acs if a not in cov]
    disagree = [a for a in acs if a in cov and set(covered.get(a, [])) != cov[a]]
    say(BAD if miss else OK, f"Coverage sheet lists {len(cov)} of {len(acs)} ACs" + (f"; missing: {', '.join(miss)}" if miss else ""))
    say(BAD if disagree else OK, "Coverage matches the test cases" +
        (f"; differs for {len(disagree)}: {', '.join(disagree[:12])}{' …' if len(disagree) > 12 else ''}" if disagree else ""))

# every test case complete, Status Not run
empty = [get(r, c_id) for r in rows if not get(r, c_steps) or not get(r, c_exp) or not get(r, c_data)]
say(BAD if empty else OK, "every test case has steps, test data and an expected result" + (f"; not: {', '.join(empty[:15])}" if empty else ""))
statuses = Counter(get(r, c_status) for r in rows)
say(OK if set(statuses) == {"Not run"} else BAD, f"Status: {dict(statuses)}")

# Won't stories: only absence checks
wont_rows = [r for r in rows if any(prio_of.get(a) == "Won't" for a in ac_re.findall(get(r, c_ac)))]
absent_words = re.compile(r"\bno\b|\bnot\b|isn't|is not|doesn't|does not|absent|only|behaves as in|same as", re.I)
odd = [get(r, c_id) for r in wont_rows if not absent_words.search(get(r, c_exp))]
say(BAD if odd else OK, f"Won't stories: {len(wont_rows)} test cases, all absence checks" + (f"; check: {', '.join(odd)}" if odd else ""))

# vague expected results
VAGUE = re.compile(r"\bcorrect(ly)?\b|\bproperly\b|\bas expected\b|an error (message )?is shown|error message is (shown|displayed)|"
                   r"\bappropriate\b|\bworks\b|\bsuccessfully\b|\bshould\b|\bvalid(ly)?\b result", re.I)
vague = [(get(r, c_id), VAGUE.search(get(r, c_exp)).group(0)) for r in rows if VAGUE.search(get(r, c_exp))]
say(BAD if vague else OK, "no vague expected results" + (f"; {len(vague)}: " + ", ".join(f"{i} ('{w}')" for i, w in vague[:12]) if vague else ""))

# ---------- Reach example only: change KEYS and MSGS for your project, or delete this block ----------
# key numbers from the sample file
alltext = "\n".join(get(r, c_exp) + " " + get(r, c_title) + " " + get(r, c_steps) for r in rows)
exp_text = "\n".join(get(r, c_exp) for r in rows)
KEYS = [("month 10.5% A+", r"10\.5\s?%.*A\+|A\+.*10\.5\s?%"), ("Instagram 11.1% A+", r"11\.1\s?%"), ("YouTube 2.6% C", r"2\.6\s?%"),
        ("row 7 at 3.0% B", r"3\.0\s?%[^.\n]*\bB\b"), ("row 2 at 745 and 6.0%", r"745"), ("row 17 at 0.8% D", r"0\.8\s?%"),
        ("top post row 18 at 55.9%", r"55\.9\s?%"), ("row 10 above row 4 (9.71% vs 9.69%)", r"9\.71|9\.69"),
        ("14 graded, 9 not graded", r"14 posts graded|14 graded")]
for name, pat in KEYS:
    say(OK if re.search(pat, exp_text) else BAD, f"expected results include {name}")
MSGS = ['Row 8: Views must be a full number, like 12400 or 12,400. "12.4K" is not accepted.',
        'Row 9: "Facebook" is not supported. Use Instagram or YouTube.',
        "Row 15: 14-09-2062 is in the future.", "Row 19: Likes is empty. Enter 0 if there were none.",
        "Row 20: The year must have 4 digits, for example 21-09-2026.",
        'Row 22: Views must be a full number, like 12400 or 12,400. "1.2L" is not accepted.', "Row 24: 31-09-2026 is not a real date."]
fold = lambda s: s.replace("“", '"').replace("”", '"').replace("’", "'")
missing_msgs = [m for m in MSGS if fold(m) not in fold(exp_text)]
say(BAD if missing_msgs else OK, "RC-04.2's seven messages appear word for word" + (f"; missing: {missing_msgs}" if missing_msgs else ""))
half = re.search(r"\b\d+\.\d5\s?%", alltext)
say(OK if half else BAD, "a case lands exactly on a rounding half (x.x5%)" + (f": {half.group(0)}" if half else "; none"))

# ---------- counts ----------
types = Counter(get(r, c_type) for r in rows)
prios = Counter(get(r, c_prio) for r in rows)
per_ac = [len(covered[a]) for a in acs if a in covered]
print(f"\ntypes: {dict(types)}\npriorities: {dict(prios)}")
if per_ac:
    print(f"test cases per AC: min {min(per_ac)}, max {max(per_ac)}, mean {sum(per_ac) / len(per_ac):.1f}")
if sheets["data"] is not None:
    print(f"test data files: {sheets['data'].max_row - 1}")
if sheets["open"] is not None:
    print(f"open questions: {sheets['open'].max_row - 1}")
# ---------- the open questions file for the BA and the devs ----------
if oq_file is not None and oq_file.exists():
    owb = openpyxl.load_workbook(oq_file)
    ows = find_sheet(owb, "Open questions", "Questions") or owb.worksheets[0]
    ohm, ohdr = header_map(ows)
    o_id, o_q = col(ohm, "ID"), col(ohm, "Question")
    o_refs = col(ohm, "Story, AC or rule IDs", "Story / AC / Rule", "Story AC or rule IDs", "IDs", "Refs", "Rule")
    o_blk = col(ohm, "Blocking or Not blocking", "Blocking")
    o_tcs = col(ohm, "Test cases affected", "Affected test cases", "Test cases")
    o_who = col(ohm, "Who answers", "Owner")
    o_ans, o_st = col(ohm, "Answer"), col(ohm, "Status")
    orows = [r for r in ows.iter_rows(min_row=2, values_only=True) if r and any(v not in (None, "") for v in r)]
    print(f"\nopen questions: '{oq_file.name}', sheet '{ows.title}' ({', '.join(owb.sheetnames)}): {len(orows)} questions")
    print(f"columns: {', '.join(str(h) for h in ohdr if h)}")
    fz = ows.freeze_panes
    say(OK if fz and fz != "A1" else BAD, f"open questions header frozen: {fz or 'no'}")
    miss = [n for n, c in [("ID", o_id), ("Question", o_q), ("Story, AC or rule IDs", o_refs), ("Blocking", o_blk),
                           ("Test cases affected", o_tcs), ("Who answers", o_who), ("Answer", o_ans), ("Status", o_st)] if c is None]
    say(BAD if miss else OK, "open questions columns present" + (f"; missing: {', '.join(miss)}" if miss else ""))
    oq_ids = [get(r, o_id) for r in orows]
    d = [i for i, n in Counter(oq_ids).items() if n > 1]
    say(BAD if d else OK, f"open question IDs unique ({oq_ids[0] if oq_ids else ''} … {oq_ids[-1] if oq_ids else ''})" + (f"; duplicates: {d}" if d else ""))
    print(f"blocking: {dict(Counter(get(r, o_blk) for r in orows))}")
    print(f"who answers: {dict(Counter(get(r, o_who) for r in orows))}")
    st = Counter(get(r, o_st) for r in orows)
    say(OK if set(st) == {"Open"} else BAD, f"open questions Status: {dict(st)}")
    answered = [get(r, o_id) for r in orows if get(r, o_ans)]
    say(BAD if answered else OK, "every Answer is empty" + (f"; filled: {', '.join(answered)}" if answered else ""))
    tc_ids = set(ids)
    bad_tc = sorted({t for r in orows for t in re.findall(r"TC-\d+", get(r, o_tcs)) if t not in tc_ids})
    say(BAD if bad_tc else OK, "open questions name only test cases that exist" + (f"; unknown: {', '.join(bad_tc)}" if bad_tc else ""))
    bad_ac = sorted({a for r in orows for a in ac_re.findall(get(r, o_refs)) if a not in acs})
    say(BAD if bad_ac else OK, "open questions name only ACs that exist" + (f"; unknown: {', '.join(bad_ac)}" if bad_ac else ""))
    q_re = re.compile(r"\bO?Q-\d+\b")
    q_in_tc = sorted({q for r in rows for v in r if v for q in q_re.findall(str(v))})
    unknown_q = [q for q in q_in_tc if q not in set(oq_ids)]
    say(BAD if unknown_q else OK, f"test cases mention {len(q_in_tc)} open questions" +
        (f"; not in the file: {', '.join(unknown_q)}" if unknown_q else ""))

print(f"\n{problems} problem(s)")
